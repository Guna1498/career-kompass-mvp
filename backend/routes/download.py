import html
import re
from io import BytesIO

from fastapi import APIRouter, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel

router = APIRouter()


class DownloadRequest(BaseModel):
    personal_details: dict = {}
    about_me: dict = {}
    work_experience: list = []
    education: list = []
    projects: list = []
    skills: dict = {}
    keywords_added: list[str] = []
    summary: str = ""
    template: str = "modern"


class CoverLetterDownloadRequest(BaseModel):
    cover_letter: str
    subject_line: str
    user_name: str


# ── Helpers ────────────────────────────────────────────────────────────────────

def _contact_line(pd: dict) -> str:
    parts = [pd.get("email"), pd.get("phone"), pd.get("location"),
             pd.get("linkedin"), pd.get("github")]
    return "  |  ".join(p for p in parts if p)


def _esc(text) -> str:
    return html.escape(str(text or ""))


def _parse_skills(text: str) -> list[tuple[str, str]]:
    """
    Split a skills blob like 'Programming: Python, SQL | Databases: PostgreSQL'
    into [(label, skills), ...] pairs so each category renders on its own line.
    Handles newline, pipe, and double-space delimiters between categories.
    Returns [] if no category structure is detected (caller falls back to plain text).
    """
    if not text:
        return []
    # Normalise pipe and 2+ spaces to newline so all formats look the same
    normalised = re.sub(r'\s*\|\s*|\s{2,}', '\n', text.strip())
    pairs = []
    for line in normalised.splitlines():
        line = line.strip()
        m = re.match(r'^([A-Za-z][A-Za-z\s&/]*):\s*(.+)$', line)
        if m:
            pairs.append((m.group(1).strip(), m.group(2).strip().rstrip(',')))
    return pairs


# ── DOCX ───────────────────────────────────────────────────────────────────────

@router.post("/download-cv/docx")
def download_docx(body: DownloadRequest):
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor

        GREEN = RGBColor(0x1D, 0x9E, 0x75)
        DARK_GREEN = RGBColor(0x1D, 0x4D, 0x3A)
        GREY = RGBColor(0x6B, 0x72, 0x80)

        doc = Document()

        for section in doc.sections:
            section.top_margin = Pt(40)
            section.bottom_margin = Pt(40)
            section.left_margin = Pt(56)
            section.right_margin = Pt(56)

        def add_run(para, text, bold=False, size=11, color=None):
            run = para.add_run(text)
            run.bold = bold
            run.font.size = Pt(size)
            if color:
                run.font.color.rgb = color
            return run

        def section_heading(label):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            add_run(p, label.upper(), bold=True, size=11, color=DARK_GREEN)
            return p

        # ── Personal header ── (Issue 1 fix: each element its own para with explicit spacing)
        pd = body.personal_details

        name_para = doc.add_paragraph()
        name_para.paragraph_format.space_before = Pt(0)
        name_para.paragraph_format.space_after = Pt(2)
        add_run(name_para, pd.get("name", ""), bold=True, size=20)

        if pd.get("title"):
            title_para = doc.add_paragraph()
            title_para.paragraph_format.space_before = Pt(0)
            title_para.paragraph_format.space_after = Pt(2)
            add_run(title_para, pd.get("title", ""), size=12, color=GREEN)

        contact = _contact_line(pd)
        if contact:
            contact_para = doc.add_paragraph()
            contact_para.paragraph_format.space_before = Pt(0)
            contact_para.paragraph_format.space_after = Pt(8)
            add_run(contact_para, contact, size=9, color=GREY)

        # ── About Me ──
        am = body.about_me
        if am.get("rewritten"):
            section_heading("About Me")
            p = doc.add_paragraph(am.get("rewritten", ""))
            p.paragraph_format.space_after = Pt(4)

        # ── Work Experience ──
        if body.work_experience:
            section_heading("Work Experience")
            for i, job in enumerate(body.work_experience):
                job_para = doc.add_paragraph()
                job_para.paragraph_format.space_before = Pt(6)
                job_para.paragraph_format.space_after = Pt(1)
                add_run(job_para, job.get("company", ""), bold=True, size=11)
                if job.get("role"):
                    add_run(job_para, f"  —  {job.get('role', '')}", size=11)

                meta_parts = [job.get("period", ""), job.get("location", "")]
                meta = "  |  ".join(p for p in meta_parts if p)
                if meta:
                    meta_para = doc.add_paragraph()
                    meta_para.paragraph_format.space_after = Pt(2)
                    add_run(meta_para, meta, size=9, color=GREY)

                for b in job.get("bullets", []):
                    bp = doc.add_paragraph(b.get("rewritten", ""), style="List Bullet")
                    bp.paragraph_format.space_after = Pt(6)

        # ── Education ──
        if body.education:
            section_heading("Education")
            for edu in body.education:
                edu_para = doc.add_paragraph()
                edu_para.paragraph_format.space_before = Pt(4)
                edu_para.paragraph_format.space_after = Pt(1)
                add_run(edu_para, edu.get("degree", ""), bold=True, size=11)
                if edu.get("institution"):
                    add_run(edu_para, f"  —  {edu.get('institution', '')}", size=11)

                meta_parts = [edu.get("period", ""), edu.get("location", "")]
                meta = "  |  ".join(p for p in meta_parts if p)
                if meta:
                    meta_para = doc.add_paragraph()
                    meta_para.paragraph_format.space_after = Pt(2)
                    add_run(meta_para, meta, size=9, color=GREY)

        # ── Projects ──
        if body.projects:
            section_heading("Projects")
            for proj in body.projects:
                proj_para = doc.add_paragraph()
                proj_para.paragraph_format.space_before = Pt(6)
                proj_para.paragraph_format.space_after = Pt(1)
                add_run(proj_para, proj.get("name", ""), bold=True, size=11)
                if proj.get("period"):
                    add_run(proj_para, f"  |  {proj.get('period', '')}", size=10, color=GREY)
                for b in proj.get("bullets", []):
                    bp = doc.add_paragraph(b.get("rewritten", ""), style="List Bullet")
                    bp.paragraph_format.space_after = Pt(6)

        # ── Skills (Issue 2 fix: each category on its own line, bold label) ──
        sk = body.skills
        if sk.get("rewritten"):
            section_heading("Skills")
            skill_pairs = _parse_skills(sk.get("rewritten", ""))
            if skill_pairs:
                for label, skills in skill_pairs:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(3)
                    add_run(p, f"{label}: ", bold=True, size=11)
                    add_run(p, skills, size=11)
            else:
                doc.add_paragraph(sk.get("rewritten", ""))

        # Keywords Added intentionally omitted (Issue 3)

        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        return Response(
            content=buffer.read(),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": "attachment; filename=rewritten_cv.docx"},
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate DOCX: {e}")


# ── PDF ────────────────────────────────────────────────────────────────────────

def _build_modern_story(body: DownloadRequest) -> list:
    """Green-accent template (original default design)."""
    from reportlab.lib.colors import HexColor
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import HRFlowable, Paragraph, Spacer

    GREEN = HexColor("#1D9E75")
    DARK = HexColor("#111827")
    GREY = HexColor("#6B7280")

    s_name    = ParagraphStyle("PKName",    fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=DARK,  spaceAfter=0)
    s_title   = ParagraphStyle("PKTitle",   fontName="Helvetica",      fontSize=12, leading=15, textColor=GREEN, spaceAfter=0)
    s_contact = ParagraphStyle("PKContact", fontName="Helvetica",      fontSize=9,  leading=12, textColor=GREY,  spaceAfter=0)
    s_section = ParagraphStyle("PKSection", fontName="Helvetica-Bold", fontSize=11, textColor=GREEN, spaceBefore=12, spaceAfter=12)
    s_job     = ParagraphStyle("PKJob",     fontName="Helvetica-Bold", fontSize=10, textColor=DARK,  spaceBefore=0,  spaceAfter=1)
    s_meta    = ParagraphStyle("PKMeta",    fontName="Helvetica",      fontSize=9,  textColor=GREY,  spaceAfter=3)
    s_body    = ParagraphStyle("PKBody",    fontName="Helvetica",      fontSize=10, textColor=DARK,  spaceAfter=4, leading=15)
    s_bullet  = ParagraphStyle("PKBullet",  fontName="Helvetica",      fontSize=10, textColor=DARK,  leftIndent=20, spaceAfter=6, leading=15)
    s_skill   = ParagraphStyle("PKSkill",   fontName="Helvetica",      fontSize=10, textColor=DARK,  spaceAfter=4,  leading=15)

    def hr():
        return HRFlowable(width="100%", thickness=0.5, color=GREEN, spaceBefore=2, spaceAfter=8)

    story = []

    pd = body.personal_details
    if pd.get("name"):
        story.append(Paragraph(_esc(pd["name"]), s_name))
        story.append(Spacer(1, 8))
    if pd.get("title"):
        story.append(Paragraph(_esc(pd["title"]), s_title))
        story.append(Spacer(1, 6))
    contact = _contact_line(pd)
    if contact:
        story.append(Paragraph(_esc(contact), s_contact))
        story.append(Spacer(1, 16))

    am = body.about_me
    if am.get("rewritten"):
        story += [Paragraph("ABOUT ME", s_section), hr(),
                  Paragraph(_esc(am["rewritten"]), s_body)]

    if body.work_experience:
        story += [Paragraph("WORK EXPERIENCE", s_section), hr()]
        for i, job in enumerate(body.work_experience):
            label = job.get("company", "")
            if job.get("role"):
                label += f"  —  {job['role']}"
            story.append(Paragraph(_esc(label), s_job))

            meta_parts = [job.get("period", ""), job.get("location", "")]
            meta = "  |  ".join(p for p in meta_parts if p)
            if meta:
                story.append(Paragraph(_esc(meta), s_meta))

            for b in job.get("bullets", []):
                story.append(Paragraph(f"• {_esc(b.get('rewritten', ''))}", s_bullet))

            if i < len(body.work_experience) - 1:
                story.append(Spacer(1, 8))

    if body.education:
        story += [Paragraph("EDUCATION", s_section), hr()]
        for edu in body.education:
            label = edu.get("degree", "")
            if edu.get("institution"):
                label += f"  —  {edu['institution']}"
            story.append(Paragraph(_esc(label), s_job))

            meta_parts = [edu.get("period", ""), edu.get("location", "")]
            meta = "  |  ".join(p for p in meta_parts if p)
            if meta:
                story.append(Paragraph(_esc(meta), s_meta))

    if body.projects:
        story += [Paragraph("PROJECTS", s_section), hr()]
        for i, proj in enumerate(body.projects):
            label = proj.get("name", "")
            if proj.get("period"):
                label += f"  |  {proj['period']}"
            story.append(Paragraph(_esc(label), s_job))
            for b in proj.get("bullets", []):
                story.append(Paragraph(f"• {_esc(b.get('rewritten', ''))}", s_bullet))

            if i < len(body.projects) - 1:
                story.append(Spacer(1, 8))

    sk = body.skills
    if sk.get("rewritten"):
        story += [Paragraph("SKILLS", s_section), hr()]
        skill_pairs = _parse_skills(sk.get("rewritten", ""))
        if skill_pairs:
            for label, skills in skill_pairs:
                story.append(Paragraph(
                    f"<b>{_esc(label)}:</b> {_esc(skills)}",
                    s_skill,
                ))
        else:
            story.append(Paragraph(_esc(sk["rewritten"]), s_body))

    return story


def _build_classic_story(body: DownloadRequest) -> list:
    """Minimal black-and-white serif template — centered header, understated rules."""
    from reportlab.lib.colors import HexColor, black
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import HRFlowable, Paragraph, Spacer

    GREY = HexColor("#595959")

    s_name    = ParagraphStyle("CLSName",    fontName="Times-Bold",   fontSize=22, leading=26, textColor=black, alignment=TA_CENTER, spaceAfter=0)
    s_title   = ParagraphStyle("CLSTitle",   fontName="Times-Italic", fontSize=12, leading=15, textColor=GREY,  alignment=TA_CENTER, spaceAfter=0)
    s_contact = ParagraphStyle("CLSContact", fontName="Times-Roman",  fontSize=9,  leading=12, textColor=GREY,  alignment=TA_CENTER, spaceAfter=0)
    s_section = ParagraphStyle("CLSSection", fontName="Times-Bold",   fontSize=11, textColor=black, spaceBefore=14, spaceAfter=6, tracking=1)
    s_job     = ParagraphStyle("CLSJob",     fontName="Times-Bold",   fontSize=10.5, textColor=black, spaceBefore=0, spaceAfter=1)
    s_meta    = ParagraphStyle("CLSMeta",    fontName="Times-Italic", fontSize=9,  textColor=GREY,  spaceAfter=3)
    s_body    = ParagraphStyle("CLSBody",    fontName="Times-Roman",  fontSize=10.5, textColor=black, spaceAfter=4, leading=15)
    s_bullet  = ParagraphStyle("CLSBullet",  fontName="Times-Roman",  fontSize=10.5, textColor=black, leftIndent=18, spaceAfter=6, leading=15)
    s_skill   = ParagraphStyle("CLSSkill",   fontName="Times-Roman",  fontSize=10.5, textColor=black, spaceAfter=4, leading=15)

    def hr():
        return HRFlowable(width="100%", thickness=0.75, color=black, spaceBefore=2, spaceAfter=10)

    story = []

    pd = body.personal_details
    if pd.get("name"):
        story.append(Paragraph(_esc(pd["name"]).upper(), s_name))
        story.append(Spacer(1, 6))
    if pd.get("title"):
        story.append(Paragraph(_esc(pd["title"]), s_title))
        story.append(Spacer(1, 6))
    contact = _contact_line(pd)
    if contact:
        story.append(Paragraph(_esc(contact), s_contact))
        story.append(Spacer(1, 14))
    story.append(hr())

    am = body.about_me
    if am.get("rewritten"):
        story += [Paragraph("ABOUT ME", s_section),
                  Paragraph(_esc(am["rewritten"]), s_body)]

    if body.work_experience:
        story.append(Paragraph("WORK EXPERIENCE", s_section))
        for i, job in enumerate(body.work_experience):
            label = job.get("company", "")
            if job.get("role"):
                label += f", {job['role']}"
            story.append(Paragraph(_esc(label), s_job))

            meta_parts = [job.get("period", ""), job.get("location", "")]
            meta = "  |  ".join(p for p in meta_parts if p)
            if meta:
                story.append(Paragraph(_esc(meta), s_meta))

            for b in job.get("bullets", []):
                story.append(Paragraph(f"– {_esc(b.get('rewritten', ''))}", s_bullet))

            if i < len(body.work_experience) - 1:
                story.append(Spacer(1, 8))

    if body.education:
        story.append(Paragraph("EDUCATION", s_section))
        for edu in body.education:
            label = edu.get("degree", "")
            if edu.get("institution"):
                label += f", {edu['institution']}"
            story.append(Paragraph(_esc(label), s_job))

            meta_parts = [edu.get("period", ""), edu.get("location", "")]
            meta = "  |  ".join(p for p in meta_parts if p)
            if meta:
                story.append(Paragraph(_esc(meta), s_meta))

    if body.projects:
        story.append(Paragraph("PROJECTS", s_section))
        for i, proj in enumerate(body.projects):
            label = proj.get("name", "")
            if proj.get("period"):
                label += f"  |  {proj['period']}"
            story.append(Paragraph(_esc(label), s_job))
            for b in proj.get("bullets", []):
                story.append(Paragraph(f"– {_esc(b.get('rewritten', ''))}", s_bullet))

            if i < len(body.projects) - 1:
                story.append(Spacer(1, 8))

    sk = body.skills
    if sk.get("rewritten"):
        story.append(Paragraph("SKILLS", s_section))
        skill_pairs = _parse_skills(sk.get("rewritten", ""))
        if skill_pairs:
            for label, skills in skill_pairs:
                story.append(Paragraph(
                    f"<b>{_esc(label)}:</b> {_esc(skills)}",
                    s_skill,
                ))
        else:
            story.append(Paragraph(_esc(sk["rewritten"]), s_body))

    return story


def _header_row(left_html: str, right_html: str, s_left, s_right, col_widths) -> "Table":
    """Borderless two-column row: bold entry on the left, italic dates/location on the right."""
    from reportlab.platypus import Paragraph, Table, TableStyle

    t = Table([[Paragraph(left_html, s_left), Paragraph(right_html, s_right)]], colWidths=col_widths)
    t.setStyle(TableStyle([
        ("ALIGN", (1, 0), (1, 0), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    return t


def _build_professional_story(body: DownloadRequest) -> list:
    """LaTeX-resume-style template: centered serif header, tight section rules,
    two-column job/education rows (bold entry left, italic dates/location right)."""
    from reportlab.lib.colors import HexColor, black
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import HRFlowable, Paragraph, Spacer

    GREY = HexColor("#595959")
    LEFT_COL, RIGHT_COL = 4.2 * inch, 2.07 * inch

    s_name    = ParagraphStyle("PFName",    fontName="Times-Bold",   fontSize=19, leading=23, textColor=black, alignment=TA_CENTER, spaceAfter=0)
    s_title   = ParagraphStyle("PFTitle",   fontName="Times-Roman",  fontSize=11, leading=14, textColor=black, alignment=TA_CENTER, spaceAfter=0)
    s_contact = ParagraphStyle("PFContact", fontName="Times-Roman",  fontSize=9,  leading=12, textColor=GREY,  alignment=TA_CENTER, spaceAfter=0)
    s_section = ParagraphStyle("PFSection", fontName="Times-Bold",   fontSize=11, textColor=black, spaceBefore=12, spaceAfter=2)
    s_left    = ParagraphStyle("PFRowLeft", fontName="Times-Bold",   fontSize=10.5, textColor=black, leading=13)
    s_right   = ParagraphStyle("PFRowRight",fontName="Times-Italic", fontSize=9.5,  textColor=black, leading=13)
    s_body    = ParagraphStyle("PFBody",    fontName="Times-Roman",  fontSize=10, textColor=black, spaceAfter=4, leading=13.5)
    s_bullet  = ParagraphStyle("PFBullet",  fontName="Times-Roman",  fontSize=10, textColor=black, leftIndent=14, spaceAfter=3, leading=13)
    s_skill   = ParagraphStyle("PFSkill",   fontName="Times-Roman",  fontSize=10, textColor=black, spaceAfter=3, leading=13)

    def hr():
        return HRFlowable(width="100%", thickness=0.75, color=black, spaceBefore=0, spaceAfter=8)

    story = []

    pd = body.personal_details
    if pd.get("name"):
        story.append(Paragraph(_esc(pd["name"]), s_name))
        story.append(Spacer(1, 4))
    if pd.get("title"):
        story.append(Paragraph(_esc(pd["title"]), s_title))
        story.append(Spacer(1, 4))
    contact = _contact_line(pd)
    if contact:
        story.append(Paragraph(_esc(contact).replace("  |  ", "  ·  "), s_contact))
    story.append(Spacer(1, 10))

    am = body.about_me
    if am.get("rewritten"):
        story += [Paragraph("ABOUT ME", s_section), hr(),
                  Paragraph(_esc(am["rewritten"]), s_body)]

    if body.work_experience:
        story += [Paragraph("WORK EXPERIENCE", s_section), hr()]
        for i, job in enumerate(body.work_experience):
            left = _esc(job.get("company", ""))
            if job.get("role"):
                left += f", {_esc(job['role'])}"
            right = "  |  ".join(_esc(p) for p in [job.get("period", ""), job.get("location", "")] if p)
            story.append(_header_row(left, right, s_left, s_right, [LEFT_COL, RIGHT_COL]))

            for b in job.get("bullets", []):
                story.append(Paragraph(f"• {_esc(b.get('rewritten', ''))}", s_bullet))

            if i < len(body.work_experience) - 1:
                story.append(Spacer(1, 6))

    if body.education:
        story += [Paragraph("EDUCATION", s_section), hr()]
        for edu in body.education:
            left = _esc(edu.get("degree", ""))
            if edu.get("institution"):
                left += f", {_esc(edu['institution'])}"
            right = "  |  ".join(_esc(p) for p in [edu.get("period", ""), edu.get("location", "")] if p)
            story.append(_header_row(left, right, s_left, s_right, [LEFT_COL, RIGHT_COL]))

    if body.projects:
        story += [Paragraph("PROJECTS", s_section), hr()]
        for i, proj in enumerate(body.projects):
            left = _esc(proj.get("name", ""))
            right = _esc(proj.get("period", ""))
            story.append(_header_row(left, right, s_left, s_right, [LEFT_COL, RIGHT_COL]))

            for b in proj.get("bullets", []):
                story.append(Paragraph(f"• {_esc(b.get('rewritten', ''))}", s_bullet))

            if i < len(body.projects) - 1:
                story.append(Spacer(1, 6))

    sk = body.skills
    if sk.get("rewritten"):
        story += [Paragraph("SKILLS", s_section), hr()]
        skill_pairs = _parse_skills(sk.get("rewritten", ""))
        if skill_pairs:
            for label, skills in skill_pairs:
                story.append(Paragraph(
                    f"<b>{_esc(label)}:</b> {_esc(skills)}",
                    s_skill,
                ))
        else:
            story.append(Paragraph(_esc(sk["rewritten"]), s_body))

    return story


_PDF_TEMPLATES = {
    "modern": _build_modern_story,
    "classic": _build_classic_story,
    "professional": _build_professional_story,
}


@router.post("/download-cv/pdf")
def download_pdf(body: DownloadRequest):
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import SimpleDocTemplate

        buffer = BytesIO()
        doc_pdf = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=72,
            leftMargin=72,
            topMargin=72,
            bottomMargin=72,
        )

        build_story = _PDF_TEMPLATES.get(body.template, _build_modern_story)
        doc_pdf.build(build_story(body))
        buffer.seek(0)

        return Response(
            content=buffer.read(),
            media_type="application/pdf",
            headers={"Content-Disposition": "attachment; filename=rewritten_cv.pdf"},
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate PDF: {e}")


# ── Cover Letter DOCX ──────────────────────────────────────────────────────────

@router.post("/download-cover-letter/docx")
def download_cover_letter_docx(body: CoverLetterDownloadRequest):
    try:
        from datetime import date
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement

        GREY = RGBColor(0x6B, 0x72, 0x80)
        GREEN = RGBColor(0x1D, 0x9E, 0x75)

        doc = Document()
        for section in doc.sections:
            section.top_margin = Pt(72)
            section.bottom_margin = Pt(72)
            section.left_margin = Pt(72)
            section.right_margin = Pt(72)

        def add_run(para, text, bold=False, size=11, color=None):
            run = para.add_run(text)
            run.bold = bold
            run.font.size = Pt(size)
            if color:
                run.font.color.rgb = color
            return run

        def add_para(text="", space_after=6):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(space_after)
            return p

        # Name
        name_p = add_para(space_after=2)
        add_run(name_p, body.user_name, bold=True, size=14)

        # Date
        date_p = add_para(space_after=2)
        add_run(date_p, date.today().strftime("%d %B %Y"), size=10, color=GREY)

        # Subject line
        subject_p = add_para(space_after=8)
        add_run(subject_p, body.subject_line, bold=True, size=11)

        # Divider (paragraph border bottom)
        divider_p = add_para(space_after=10)
        pPr = divider_p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "4")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "1D9E75")
        pBdr.append(bottom)
        pPr.append(pBdr)

        # Cover letter body — split into paragraphs
        paragraphs = [p.strip() for p in body.cover_letter.split("\n\n") if p.strip()]
        for para_text in paragraphs:
            p = add_para(space_after=6)
            add_run(p, para_text, size=11)

        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        return Response(
            content=buffer.read(),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": "attachment; filename=cover_letter.docx"},
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate cover letter DOCX: {e}")


# ── Cover Letter PDF ───────────────────────────────────────────────────────────

@router.post("/download-cover-letter/pdf")
def download_cover_letter_pdf(body: CoverLetterDownloadRequest):
    try:
        from datetime import date
        from reportlab.lib.colors import HexColor, black
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

        buffer = BytesIO()
        doc_pdf = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=72,
            leftMargin=72,
            topMargin=72,
            bottomMargin=72,
        )

        GREEN = HexColor("#1D9E75")
        DARK = HexColor("#111827")
        GREY = HexColor("#6B7280")

        s_name    = ParagraphStyle("CLName",    fontName="Helvetica-Bold", fontSize=16, textColor=DARK,  spaceAfter=0)
        s_date    = ParagraphStyle("CLDate",    fontName="Helvetica",      fontSize=9,  textColor=GREY,  spaceAfter=0)
        s_subject = ParagraphStyle("CLSubject", fontName="Helvetica-Bold", fontSize=11, textColor=GREEN, spaceAfter=0)
        s_body    = ParagraphStyle("CLBody",    fontName="Helvetica",      fontSize=10, textColor=DARK,  spaceAfter=0, leading=16)

        story = [
            Paragraph(_esc(body.user_name), s_name),
            Spacer(1, 4),
            Paragraph(_esc(date.today().strftime("%d %B %Y")), s_date),
            Spacer(1, 6),
            Paragraph(_esc(body.subject_line), s_subject),
            Spacer(1, 8),
            HRFlowable(width="100%", thickness=0.5, color=GREEN, spaceBefore=0, spaceAfter=12),
        ]

        paragraphs = [p.strip() for p in body.cover_letter.split("\n\n") if p.strip()]
        for para_text in paragraphs:
            story.append(Paragraph(_esc(para_text), s_body))
            story.append(Spacer(1, 10))

        doc_pdf.build(story)
        buffer.seek(0)

        return Response(
            content=buffer.read(),
            media_type="application/pdf",
            headers={"Content-Disposition": "attachment; filename=cover_letter.pdf"},
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate cover letter PDF: {e}")
